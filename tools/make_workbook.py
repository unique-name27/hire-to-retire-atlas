#!/usr/bin/env python3
"""Build the Project Atlas plan workbook (Excel) from the built page data.

usage: python3 build.py && python3 tools/make_workbook.py [out.xlsx]
Sheets: Read me, Dashboard, Plan (all 438 items), Milestones, Risks, Issues, Assumptions, Decisions,
RACI, Approval chains, Country gates, Comms, Training. Status, owner, go-live override and notes on the
Plan sheet are the inputs; everything on the Dashboard is a formula over the Plan sheet.
"""
import datetime as dt, json, os, re, sys
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "dist", "Project_Atlas_Plan.xlsx")
html = open(os.path.join(ROOT, "dist", "atlas.html")).read()
D = json.loads(re.search(r'<script[^>]*id="lib-data"[^>]*>(.*?)</script>', html, re.S).group(1).replace("<\\/", "</"))
PRJ = D["project"]
POL = {p["id"]: p for p in D["policies"]}
DOMAIN = {"Ethics & Governance": "Governance", "Workplace Conduct": "Relations", "Hiring & Employment": "Recruit", "Pay & Rewards": "Reward",
          "Time Off & Benefits": "Leave", "Work Arrangements": "Move", "Performance & Development": "Grow", "Data, IP & Technology": "Data",
          "Health & Safety": "Safety", "Separation": "Offboard", "Inclusion & Accessibility": "Relations"}
TIER = {"legal": "Legally required", "public": "Public-company driven", "semi": "Semiconductor-specific", "best": "Best practice"}


def ptier(p):
    t = [POL[x]["tier"] for x in p.get("policies", []) if x in POL]
    return "legal" if "legal" in t else "public" if "public" in t else "semi" if "semi" in t else "best"


items = [(p["id"], p["name"], "Policy", DOMAIN.get(p["category"], "Governance"), p["tier"], p["owner"]) for p in D["policies"]]
items += [(p["id"], p["name"], "Process", p["stage"], ptier(p), p["owner"]) for p in D["processes"]]
items += [(r["id"], r["name"], r["type"], r["domain"], r["tier"], r["owner"]) for r in D["register"]]
items.sort(key=lambda x: (PRJ["plan"][x[0]][7], x[0]))
d = lambda s: dt.date.fromisoformat(s) if s else None

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", fgColor="1E3A6E")
B = Font(name=F, size=10)
BB = Font(name=F, size=10, bold=True)
INPUT = PatternFill("solid", fgColor="FFF6C8")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
DATE = "d mmm yyyy"
wb = Workbook()


def sheet(title, headers, rows, widths, wrap_cols=(), date_cols=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c); cell.font = H; cell.fill = HFILL; cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = B; cell.border = BORDER
            cell.alignment = WRAP if cell.column in wrap_cols else TOP
            if cell.column in date_cols: cell.number_format = DATE
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2" if title == "Plan" else "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.row_dimensions[1].height = 30
    return ws


# ---------------- Read me ----------------
ws = wb.active; ws.title = "Read me"
meta = PRJ["meta"]
lines = [
    (meta["name"] + ": integrated plan workbook", Font(name=F, size=16, bold=True)),
    (meta["title"], Font(name=F, size=11, color="666666")),
    ("", B),
    (f"Program dates: {meta['start']} to {meta['end']}. Sponsor: {meta['sponsor']}. Program manager: {meta['pm']}. Plan version {meta['planVersion']}.", B),
    (f"Exported from the Hire-to-Retire Atlas on {dt.date.today().isoformat()}. The Atlas is the live source of truth for status and approvals; this workbook is a snapshot for planning and reporting.", B),
    ("", B),
    ("How to use it", BB),
    ("1. Plan sheet: one row per item (57 policies, 97 processes, 284 legal obligations). Yellow columns are inputs: Status, Owner, Go-live override and Notes.", B),
    ("2. Go live (current) uses the override when one is entered, otherwise the planned date. Overdue and Days late are formulas against today's date.", B),
    ("3. Dashboard sheet: every number is a formula over the Plan sheet, so it updates as you change statuses.", B),
    ("4. Status values: Not started, Drafting, In review, Approved, Live, Not applicable. Approved, Live and Not applicable count as done.", B),
    ("5. Stage dates come from the plan: discover 2 weeks, draft 3, review 2, approve 2, go live 1, skipping the year-end freeze (Dec 21, 2026 to Jan 1, 2027).", B),
    ("", B),
    ("Waves", BB),
] + [(f"{w['name']}: {w['start']} to {w['end']}. {w['goal']}", B) for w in PRJ["waves"]] + [
    ("", B), ("Example row (Plan sheet)", BB),
    ("POL-09 Anti-Harassment · Policy · Wave 1 · Governance, Ethics and Employee Relations · go live 29 Jan 2027 · Status: Drafting · Owner: Head of Employee Relations · Notes: Current doc on ER SharePoint; Israel bylaw and India IC addendum missing.", B),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(i, 1, t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 140

# ---------------- Plan ----------------
WS = [w["name"] for w in PRJ["workstreams"]]
STATUS = ["Not started", "Drafting", "In review", "Approved", "Live", "Not applicable"]
hdr = ["ID", "Item", "Type", "Domain", "Tier", "Workstream", "Wave", "Start", "Discover done", "Draft done", "Review done", "Approved by", "Go live (plan)",
       "Go live (override)", "Go live (current)", "Status", "Owner", "Notes", "Overdue", "Days late"]
rows = []
for iid, name, typ, dom, tier, owner in items:
    p = PRJ["plan"][iid]
    rows.append([iid, name, typ, dom, TIER[tier], WS[p[1]], PRJ["waves"][p[0]]["name"].split(":")[0], d(p[2]), d(p[3]), d(p[4]), d(p[5]), d(p[6]), d(p[7]), None, None, "Not started", owner, None, None, None])
ws = sheet("Plan", hdr, rows, [11, 42, 11, 12, 20, 34, 9, 12, 12, 12, 12, 12, 13, 13, 13, 13, 34, 40, 10, 9], wrap_cols=(2, 6, 17, 18), date_cols=(8, 9, 10, 11, 12, 13, 14, 15))
n = ws.max_row
for r in range(2, n + 1):
    ws[f"O{r}"] = f'=IF(N{r}<>"",N{r},M{r})'
    ws[f"O{r}"].number_format = DATE
    ws[f"S{r}"] = f'=IF(AND(O{r}<TODAY(),P{r}<>"Approved",P{r}<>"Live",P{r}<>"Not applicable"),"Overdue","")'
    ws[f"T{r}"] = f'=IF(S{r}="Overdue",TODAY()-O{r},"")'
    for c in ("N", "P", "Q", "R"): ws[f"{c}{r}"].fill = INPUT
dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=False); ws.add_data_validation(dv); dv.add(f"P2:P{n}")
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True); ws.add_data_validation(dvd); dvd.add(f"N2:N{n}")
red = PatternFill("solid", fgColor="FDE2E2")
ws.conditional_formatting.add(f"A2:T{n}", FormulaRule(formula=[f'$S2="Overdue"'], fill=red))
for val, col in [("Live", "1F6B43"), ("Approved", "3C8C5E"), ("In review", "6F8FC7"), ("Drafting", "D9A441")]:
    ws.conditional_formatting.add(f"P2:P{n}", CellIsRule(operator="equal", formula=[f'"{val}"'], font=Font(name=F, bold=True, color=col)))
PLAN_N = n

# ---------------- Dashboard ----------------
ds = wb.create_sheet("Dashboard", 1)
ds["A1"] = "Dashboard"; ds["A1"].font = Font(name=F, size=16, bold=True)
ds["A2"] = "All figures are formulas over the Plan sheet."; ds["A2"].font = Font(name=F, size=10, color="666666")
rng = lambda col: f"Plan!${col}$2:${col}${PLAN_N}"
ds["A4"] = "Items"; ds["B4"] = f"=COUNTA({rng('A')})"
ds["A5"] = "Done (approved, live or not applicable)"; ds["B5"] = f'=COUNTIF({rng("P")},"Approved")+COUNTIF({rng("P")},"Live")+COUNTIF({rng("P")},"Not applicable")'
ds["A6"] = "Percent done"; ds["B6"] = "=IF(B4=0,0,B5/B4)"; ds["B6"].number_format = "0.0%"
ds["A7"] = "Overdue"; ds["B7"] = f'=COUNTIF({rng("S")},"Overdue")'
ds["A8"] = "Due so far"; ds["B8"] = f'=COUNTIF({rng("O")},"<"&TODAY())'
ds["A9"] = "Overdue as share of due"; ds["B9"] = "=IF(B8=0,0,B7/B8)"; ds["B9"].number_format = "0.0%"
ds["A10"] = "Calculated rating"; ds["B10"] = '=IF(B9>=0.15,"Red",IF(B9>=0.05,"Amber","Green"))'
ds["A11"] = "Rating thresholds"; ds["B11"] = "Red at 15% or more of due items overdue, Amber at 5%. Set by the PMO; change here if the steering committee agrees different ones."
for r in range(4, 12): ds[f"A{r}"].font = BB; ds[f"B{r}"].font = B
ds["B11"].alignment = Alignment(wrap_text=True)

def grid(top, title, labels, key_col):
    ds.cell(top, 1, title).font = Font(name=F, size=12, bold=True)
    heads = [key_col[0]] + STATUS + ["Total", "Done", "% done", "Overdue"]
    for j, hh in enumerate(heads, 1):
        c = ds.cell(top + 1, j, hh); c.font = H; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, lab in enumerate(labels):
        r = top + 2 + i
        ds.cell(r, 1, lab).font = B
        for j, s in enumerate(STATUS, 2):
            ds.cell(r, j, f'=COUNTIFS({rng(key_col[1])},$A{r},{rng("P")},"{s}")').font = B
        L = get_column_letter
        ds.cell(r, 8, f"=SUM(B{r}:G{r})").font = BB
        ds.cell(r, 9, f"=E{r}+F{r}+G{r}").font = B
        c = ds.cell(r, 10, f"=IF(H{r}=0,0,I{r}/H{r})"); c.font = B; c.number_format = "0.0%"
        ds.cell(r, 11, f'=COUNTIFS({rng(key_col[1])},$A{r},{rng("S")},"Overdue")').font = B
    r = top + 2 + len(labels)
    ds.cell(r, 1, "Total").font = BB
    for j in range(2, 12):
        col = get_column_letter(j)
        c = ds.cell(r, j, f"=SUM({col}{top + 2}:{col}{r - 1})" if j != 10 else f"=IF(H{r}=0,0,I{r}/H{r})"); c.font = BB
        if j == 10: c.number_format = "0.0%"
    return r + 2

nxt = grid(13, "By wave", [w["name"].split(":")[0] for w in PRJ["waves"][1:]], ("Wave", "G"))
nxt = grid(nxt, "By workstream", WS, ("Workstream", "F"))
grid(nxt, "By type", ["Policy", "Process", "Filing", "Control", "Program", "Notice", "Record", "Committee", "Training"], ("Type", "C"))
ds.column_dimensions["A"].width = 44
for j in range(2, 12): ds.column_dimensions[get_column_letter(j)].width = 12

# ---------------- other sheets ----------------
sheet("Milestones", ["ID", "Date", "Milestone", "Wave", "Type", "Done when"], [[m["id"], d(m["date"]), m["name"], m["wave"], m["type"], m["criteria"]] for m in PRJ["milestones"]], [7, 13, 40, 7, 12, 80], wrap_cols=(3, 6), date_cols=(2,))
R = PRJ["raid"]
rs = sheet("Risks", ["ID", "Risk", "Detail", "Likelihood", "Impact", "Score", "Owner", "Mitigation", "Trigger", "Status"],
           [[x["id"], x["title"], x["detail"], x["likelihood"], x["impact"], None, x["owner"], x["mitigation"], x["trigger"], x["status"]] for x in R["risks"]],
           [7, 34, 60, 11, 10, 8, 24, 60, 40, 11], wrap_cols=(2, 3, 7, 8, 9))
for r in range(2, rs.max_row + 1):
    rs[f"F{r}"] = f'=(MATCH(D{r},{{"Low","Medium","High"}},0))*(MATCH(E{r},{{"Low","Medium","High"}},0))'
    for c in ("D", "E", "J"): rs[f"{c}{r}"].fill = INPUT
lv = DataValidation(type="list", formula1='"High,Medium,Low"'); rs.add_data_validation(lv); lv.add(f"D2:E{rs.max_row}")
sv = DataValidation(type="list", formula1='"Open,Mitigating,Closed"'); rs.add_data_validation(sv); sv.add(f"J2:J{rs.max_row}")
rs.conditional_formatting.add(f"F2:F{rs.max_row}", CellIsRule(operator="greaterThanOrEqual", formula=["6"], fill=PatternFill("solid", fgColor="F8C9C4")))
rs.cell(rs.max_row + 2, 1, "Score = likelihood (1 low, 2 medium, 3 high) x impact (same scale). 6 or more is a high risk on the dashboard and status report.").font = Font(name=F, size=9, italic=True)
sheet("Issues", ["ID", "Issue", "Detail", "Priority", "Owner", "Action", "Due", "Status"], [[x["id"], x["title"], x["detail"], x["priority"], x["owner"], x["action"], d(x["due"]), x["status"]] for x in R["issues"]], [7, 34, 60, 9, 24, 60, 12, 10], wrap_cols=(2, 3, 5, 6), date_cols=(7,))
sheet("Assumptions", ["ID", "Assumption", "Owner", "Validate by"], [[x["id"], x["text"], x["owner"], d(x["validateBy"])] for x in R["assumptions"]], [7, 90, 28, 13], wrap_cols=(2, 3), date_cols=(4,))
sheet("Decisions", ["ID", "Decision", "Detail", "Owner", "Status", "Due"], [[x["id"], x["title"], x["detail"], x["owner"], x["status"], d(x["due"])] for x in R["decisions"]], [7, 36, 70, 26, 10, 12], wrap_cols=(2, 3, 4), date_cols=(6,))
g = PRJ["governance"]
sheet("RACI", ["Activity"] + g["raci"]["roles"], [[row["activity"]] + [row["cells"].get(r, "") for r in g["raci"]["roles"]] for row in g["raci"]["rows"]], [52] + [12] * len(g["raci"]["roles"]), wrap_cols=(1,))
crow = []
for c in PRJ["chains"]:
    for i, s in enumerate(c["steps"], 1):
        crow.append([c["name"], c["appliesTo"], i, s["name"], s["role"], "All" if s["rule"] == "all" else "Any", "Yes" if s["optional"] else "No", s["guidance"], None])
cs = sheet("Approval chains", ["Chain", "Applies to", "Step", "Step name", "Role", "Must sign", "Optional", "Guidance", "Named approvers"], crow, [26, 30, 6, 30, 30, 9, 9, 60, 30], wrap_cols=(2, 4, 5, 8, 9))
for r in range(2, cs.max_row + 1): cs[f"I{r}"].fill = INPUT
JN = {c["id"]: c["name"] for c in D["countries"]}
sheet("Country gates", ["Jurisdiction", "Consultation", "Filing or registration", "Language", "Lead time", "Notes", "Verify"], [[JN.get(x["jur"], x["jur"]), x["consultation"], x["filing"], x["language"], x["leadTime"], x.get("notes", ""), "Yes" if x.get("verify") else ""] for x in PRJ["countryGates"]], [26, 50, 50, 36, 18, 50, 7], wrap_cols=(1, 2, 3, 4, 5, 6))
sheet("Comms", ["ID", "Date", "Wave", "Type", "Audience", "Message", "Channel", "Owner"], [[x["id"], d(x["date"]), x["wave"], x["type"], x["audience"], x["message"], x["channel"], x["owner"]] for x in PRJ["comms"]], [6, 12, 6, 10, 30, 70, 24, 26], wrap_cols=(5, 6, 7, 8), date_cols=(2,))
T = PRJ["training"]
sheet("Training", ["ID", "Module", "Audience", "Format", "Length", "Objectives", "Video"], [[m["id"], m["title"], m["audience"], m["format"], m["length"], "\n".join(m["objectives"]), m.get("video") or ""] for m in T["modules"]], [6, 36, 34, 26, 12, 70, 7], wrap_cols=(2, 3, 4, 6))
for w in wb.worksheets: w.sheet_view.showGridLines = w.title in ("Plan",)
wb.save(OUT)
print("wrote", OUT, "plan rows", PLAN_N - 1)
